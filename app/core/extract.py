"""Turn an uploaded file into plain text.

PDF and DOCX support are optional at runtime: if the library is missing we raise
a clear, actionable error instead of crashing with an ImportError traceback.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Callable, Dict

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx", ".csv", ".json"}


class ExtractionError(RuntimeError):
    """Raised when a document cannot be converted to text."""


def _extract_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode file as text in utf-8, utf-16, or latin-1.")


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - depends on install extras
        raise ExtractionError(
            "PDF support requires 'pypdf'. Install it with: pip install pypdf"
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF: {exc}") from exc

    pages = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append("[page {}]\n{}".format(index, text))

    if not pages:
        raise ExtractionError(
            "No selectable text found in this PDF. It is probably a scan; "
            "run OCR on it first or upload a text-based version."
        )
    return "\n\n".join(pages)


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - depends on install extras
        raise ExtractionError(
            "DOCX support requires 'python-docx'. Install it with: "
            "pip install python-docx"
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    if not parts:
        raise ExtractionError("The DOCX file contains no readable text.")
    return "\n".join(parts)


def _extract_csv(data: bytes) -> str:
    text = _extract_text(data)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise ExtractionError("The CSV file is empty.")

    header = rows[0]
    lines = []
    for row in rows[1:]:
        pairs = [
            "{}: {}".format(header[i] if i < len(header) else "col{}".format(i), value)
            for i, value in enumerate(row)
            if value.strip()
        ]
        if pairs:
            lines.append("; ".join(pairs))
    # Keep the header so the model knows the column vocabulary.
    return "Columns: {}\n\n{}".format(", ".join(header), "\n".join(lines))


def _extract_json(data: bytes) -> str:
    text = _extract_text(data)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ExtractionError("Invalid JSON: {}".format(exc)) from exc
    return json.dumps(parsed, indent=2, ensure_ascii=False)


_EXTRACTORS: Dict[str, Callable[[bytes], str]] = {
    ".txt": _extract_text,
    ".md": _extract_text,
    ".markdown": _extract_text,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".csv": _extract_csv,
    ".json": _extract_json,
}


def extract(filename: str, data: bytes) -> str:
    """Extract plain text from ``data``, dispatching on the file extension."""
    if not data:
        raise ExtractionError("The file is empty.")

    suffix = Path(filename).suffix.lower()
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        raise ExtractionError(
            "Unsupported file type '{}'. Supported: {}".format(
                suffix or "(none)", ", ".join(sorted(SUPPORTED_EXTENSIONS))
            )
        )

    text = extractor(data)
    if not text.strip():
        raise ExtractionError("No readable text could be extracted from this file.")
    return text
